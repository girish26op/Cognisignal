"""
Fill PBL PPTX template with CogniSignal project content from the DOCX report.
Preserves all template formatting, fonts, colors, and alignment.
"""

import copy
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
import datetime

SRC_PPTX = r"C:/Users/Ramesh/.gemini/antigravity/brain/13cace39-523c-4712-887d-0f14c0c0dc5f/.user_uploaded/media_1791392077635.pptx"
OUT_PPTX = r"C:\Users\Ramesh\Desktop\ml pbl\CogniSignal_PBL_Presentation.pptx"

TODAY = "07/10/2026"

# ─────────────────────────────────────────────
# Helper: replace text in a run while keeping its format
# ─────────────────────────────────────────────
def replace_run_text(run, new_text):
    run.text = new_text

# Helper: set ALL paragraph text at once (clears extra runs, keeps first run format)
def set_para_text(para, new_text):
    if not para.runs:
        run = para.add_run()
        run.text = new_text
        return
    para.runs[0].text = new_text
    for run in para.runs[1:]:
        run.text = ""

# Helper: set all paragraphs in a text frame to given lines
# lines = list of strings; each is a paragraph line
def set_tf_lines(tf, lines, start_para=0):
    """
    Overwrite paragraph text starting at start_para, using template paragraph style.
    If more lines than paragraphs exist, clone the last paragraph's XML.
    """
    from lxml import etree
    paras = tf.paragraphs

    # Ensure we have enough paragraphs
    while len(tf.paragraphs) < start_para + len(lines):
        # Clone the last paragraph element
        last_p = tf.paragraphs[-1]._p
        new_p = copy.deepcopy(last_p)
        last_p.addnext(new_p)

    for i, line in enumerate(lines):
        para = tf.paragraphs[start_para + i]
        set_para_text(para, line)

# ─────────────────────────────────────────────
# Load presentation
# ─────────────────────────────────────────────
prs = Presentation(SRC_PPTX)
slides = prs.slides

# Helper to find a shape by name (partial match)
def find_shape(slide, name_fragment):
    for shape in slide.shapes:
        if name_fragment.lower() in shape.name.lower():
            return shape
    return None

# Helper to find all Content Placeholder shapes
def find_content_placeholders(slide):
    return [s for s in slide.shapes if s.has_text_frame and "Content Placeholder" in s.name]

def find_subtitle_shapes(slide):
    return [s for s in slide.shapes if s.has_text_frame and "Subtitle" in s.name]

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 1 — Title Slide
# ─────────────────────────────────────────────────────────────────────────────
slide1 = slides[0]

# Title textbox "TITLE OF THE PROJECT"
for shape in slide1.shapes:
    if shape.has_text_frame and "TextBox" in shape.name:
        tf = shape.text_frame
        set_para_text(tf.paragraphs[0],
            "CogniSignal: An Adaptive Multi-Model ML Pipeline for\n"
            "Real-Time PCU-Weighted Dynamic Traffic Signal Optimization")
        break

# Subtitle — Students
subtitles = find_subtitle_shapes(slide1)
if subtitles:
    tf = subtitles[0].text_frame
    lines = [
        "Presentation by",
        "MITHUN PRASATH LOGESH (210425244033)",
        "GIRISH.R (210425244045)",
        "Computer Science and Business Systems",
    ]
    for i, line in enumerate(lines):
        if i < len(tf.paragraphs):
            set_para_text(tf.paragraphs[i], line)

# Subtitle — Mentor (second subtitle box)
if len(subtitles) > 1:
    tf = subtitles[1].text_frame
    lines = [
        "Mentor",
        "Dr. N. Kirubakaran B.E M.Tech",
        "Associate Professor",
        "Dept. of Computer Science and Business Systems",
    ]
    for i, line in enumerate(lines):
        if i < len(tf.paragraphs):
            set_para_text(tf.paragraphs[i], line)

# Date
for shape in slide1.shapes:
    if shape.has_text_frame and "Date" in shape.name:
        set_para_text(shape.text_frame.paragraphs[0], f"Date: {TODAY}")

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 2 — Problem & Objectives
# ─────────────────────────────────────────────────────────────────────────────
slide2 = slides[1]
content_placeholders2 = find_content_placeholders(slide2)

if len(content_placeholders2) >= 1:
    tf = content_placeholders2[0].text_frame
    # First para is the heading — keep it
    problem_lines = [
        "Problem Statement",
        "Traditional fixed-timer traffic signals operate on static schedules,",
        "regardless of real-time vehicular volume — causing severe urban gridlock,",
        "excess fuel consumption, and delayed emergency responses.",
        "Existing systems cannot differentiate vehicle type or dynamically adapt",
        "signal duration to actual road load.",
    ]
    set_tf_lines(tf, problem_lines, start_para=0)

if len(content_placeholders2) >= 2:
    tf = content_placeholders2[1].text_frame
    obj_lines = [
        "Objectives",
        "• Pre-process multi-format traffic datasets (YOLO txt, Pascal VOC xml).",
        "• Build & evaluate a Custom PyTorch CNN, Faster R-CNN, and YOLOv8.",
        "• Compute PCU density (Car=1.0, Bus=2.5, Truck=3.0, Bike=0.5, Auto=1.2).",
        "• Benchmark models: Precision, Recall, F1, mAP@0.5, FPS.",
        "• Deploy an interactive Streamlit dashboard for real-time simulation.",
    ]
    set_tf_lines(tf, obj_lines, start_para=0)

if len(content_placeholders2) >= 3:
    tf = content_placeholders2[2].text_frame
    outcome_lines = [
        "Expected Outcome",
        "An ML-based system capable of detecting multi-class vehicles,",
        "computing PCU traffic load, and dynamically recommending green-light",
        "durations (15 s – 90 s) to minimize intersection delays.",
    ]
    set_tf_lines(tf, outcome_lines, start_para=0)

# Date
for shape in slide2.shapes:
    if shape.has_text_frame and "Date" in shape.name:
        set_para_text(shape.text_frame.paragraphs[0], f"Date: {TODAY}")

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 3 — Input, Analysis & Insights
# ─────────────────────────────────────────────────────────────────────────────
slide3 = slides[2]
content_placeholders3 = find_content_placeholders(slide3)

if len(content_placeholders3) >= 1:
    tf = content_placeholders3[0].text_frame
    input_lines = [
        "Input / Data Sources",
        "• 9 GB annotated traffic image dataset (YOLO + Pascal VOC XML formats).",
        "• 5 vehicle classes: Car, Bus, Truck, Motorcycle, Auto-Rickshaw.",
        "• Tabular CSV feature data for signal-duration regression.",
        "• Standard RGB traffic camera images (640×640 / 128×128).",
    ]
    set_tf_lines(tf, input_lines, start_para=0)

if len(content_placeholders3) >= 2:
    tf = content_placeholders3[1].text_frame
    analysis_lines = [
        "Analysis / Processing",
        "• OpenCV: Grayscale conversion, edge detection, morphological ops.",
        "• Lazy batch loading prevents RAM overflow on large datasets.",
        "• 70% Train / 15% Validation / 15% Test split.",
        "• AdamW optimizer (lr=0.001) + StepLR scheduler.",
        "• Multi-task loss = CrossEntropy + 2.0 × MSE.",
    ]
    set_tf_lines(tf, analysis_lines, start_para=0)

if len(content_placeholders3) >= 3:
    tf = content_placeholders3[2].text_frame
    insights_lines = [
        "Key Insights",
        "• YOLOv8 achieves best accuracy (mAP@0.5: 0.908) and speed (65 FPS).",
        "• PCU weighting captures heavy-vehicle road load more accurately than raw count.",
        "• Random Forest tabular model: 94.2% accuracy, MAE 1.82 s.",
        "• Multi-class ROC-AUC scores all exceeded 0.94.",
    ]
    set_tf_lines(tf, insights_lines, start_para=0)

for shape in slide3.shapes:
    if shape.has_text_frame and "Date" in shape.name:
        set_para_text(shape.text_frame.paragraphs[0], f"Date: {TODAY}")

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 4 — Technical Approach
# ─────────────────────────────────────────────────────────────────────────────
slide4 = slides[3]
content_placeholders4 = find_content_placeholders(slide4)

if len(content_placeholders4) >= 1:
    tf = content_placeholders4[0].text_frame
    arch_lines = [
        "Architecture",
        "4-stage end-to-end pipeline:",
        "1. Camera Image Input",
        "2. OpenCV Feature Extraction & AI Object Detection",
        "3. PCU Traffic Load Engine",
        "4. Dynamic Signal Timer & Streamlit UI",
    ]
    set_tf_lines(tf, arch_lines, start_para=0)

if len(content_placeholders4) >= 2:
    tf = content_placeholders4[1].text_frame
    ml_lines = [
        "ML Methodology & Implementation",
        "Iter 1 — Custom PyTorch CNN from scratch (4 Conv+BN+ReLU+MaxPool blocks, dual head).",
        "Iter 2 — Faster R-CNN (MobileNetV3 FPN): Precision 0.881, 18.2 FPS.",
        "Final   — YOLOv8 Small: mAP@0.5 0.908, 65 FPS. Selected winner.",
        "Tabular — Scikit-Learn Random Forest for congestion & signal regression.",
    ]
    set_tf_lines(tf, ml_lines, start_para=0)

if len(content_placeholders4) >= 3:
    tf = content_placeholders4[2].text_frame
    tech_lines = [
        "Technology Stack",
        "Languages: Python 3.10+",
        "Frameworks: PyTorch, Torchvision, Ultralytics YOLOv8, OpenCV",
        "ML/DL Tools: Scikit-Learn, Pandas, NumPy, Seaborn",
        "UI: Streamlit (7-tab interactive dashboard)",
        "Hardware: Intel i5/i7, 16 GB RAM, NVIDIA RTX GPU (optional)",
    ]
    set_tf_lines(tf, tech_lines, start_para=0)

for shape in slide4.shapes:
    if shape.has_text_frame and "Date" in shape.name:
        set_para_text(shape.text_frame.paragraphs[0], f"Date: {TODAY}")

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 5 — Feasibility, Risk and Challenges
# ─────────────────────────────────────────────────────────────────────────────
slide5 = slides[4]
content_placeholders5 = find_content_placeholders(slide5)

if len(content_placeholders5) >= 1:
    tf = content_placeholders5[0].text_frame
    feas_lines = [
        "Feasibility",
        "• Open-source ML stack (PyTorch, YOLOv8, Scikit-Learn) — zero licensing cost.",
        "• Modular decoupled pipeline (detection, PCU, signal, UI) — easy to upgrade.",
        "• Transfer learning (pre-trained Faster R-CNN & YOLOv8) — rapid convergence.",
        "• Lazy batch loading supports 9 GB dataset on consumer laptops.",
        "• Works with standard RGB CCTV — no road induction sensors needed.",
    ]
    set_tf_lines(tf, feas_lines, start_para=0)

if len(content_placeholders5) >= 2:
    tf = content_placeholders5[1].text_frame
    risk_lines = [
        "Risk and Challenges",
        "• Image quality degrades in dense fog, heavy rain, or night without IR.",
        "• Severe visual occlusion (motorcycle hidden behind truck) causes missed counts.",
        "• Static RGB cameras may miss thermal/infrared vehicle signatures.",
        "• Signal timing is simulation-based, not physically wired to real hardware.",
        "• Single-intersection scope; no cooperative multi-intersection logic yet.",
    ]
    set_tf_lines(tf, risk_lines, start_para=0)

if len(content_placeholders5) >= 3:
    tf = content_placeholders5[2].text_frame
    mitigation_lines = [
        "Mitigation",
        "• Infrared/thermal cameras can be added as future hardware upgrade.",
        "• Ensemble detection + tracking reduces occlusion-related missed counts.",
        "• Arduino/Raspberry Pi relay integration planned for physical deployment.",
        "• Streamlit dashboard provides visual fallback for operator override.",
    ]
    set_tf_lines(tf, mitigation_lines, start_para=0)

for shape in slide5.shapes:
    if shape.has_text_frame and "Date" in shape.name:
        set_para_text(shape.text_frame.paragraphs[0], f"Date: {TODAY}")

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 6 — Result and Applications
# ─────────────────────────────────────────────────────────────────────────────
slide6 = slides[5]
content_placeholders6 = find_content_placeholders(slide6)

if len(content_placeholders6) >= 1:
    tf = content_placeholders6[0].text_frame
    result_lines = [
        "Project Output — Model Benchmark Results",
        "Model                | Precision | Recall | F1    | mAP@0.5 | FPS",
        "Custom CNN (Iter 1)  |   0.742   |  0.698 | 0.719 |  0.712  | 42.5",
        "Faster R-CNN (Iter 2)|   0.881   |  0.854 | 0.867 |  0.862  | 18.2",
        "YOLOv8 Small (Final) |   0.915   |  0.892 | 0.903 |  0.908  | 65.0",
        "Random Forest Tabular: 94.2% Accuracy | MAE 1.82 s | ROC-AUC > 0.94",
    ]
    set_tf_lines(tf, result_lines, start_para=0)

if len(content_placeholders6) >= 2:
    tf = content_placeholders6[1].text_frame
    app_lines = [
        "Applications",
        "• Smart city traffic signal control at urban intersections.",
        "• Emergency vehicle priority preemption & clearance corridors.",
        "• Real-time traffic density monitoring via CCTV analytics.",
        "• Transport authority dashboards for congestion reporting.",
        "• Integration with GPS navigation for adaptive routing.",
    ]
    set_tf_lines(tf, app_lines, start_para=0)

for shape in slide6.shapes:
    if shape.has_text_frame and "Date" in shape.name:
        set_para_text(shape.text_frame.paragraphs[0], f"Date: {TODAY}")

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 7 — Conclusion, Limitations & Future Scope
# ─────────────────────────────────────────────────────────────────────────────
slide7 = slides[6]
content_placeholders7 = find_content_placeholders(slide7)

if len(content_placeholders7) >= 1:
    tf = content_placeholders7[0].text_frame
    conc_lines = [
        "Conclusion",
        "CogniSignal successfully demonstrates an adaptive multi-model ML pipeline",
        "for real-time traffic signal optimization. By combining PyTorch deep learning",
        "object detection, PCU load weighting, and tabular Random Forest models, the",
        "system dynamically adjusts green-light durations (15 s – 90 s) based on live",
        "vehicle density — reducing congestion without costly hardware upgrades.",
    ]
    set_tf_lines(tf, conc_lines, start_para=0)

if len(content_placeholders7) >= 2:
    tf = content_placeholders7[1].text_frame
    future_lines = [
        "Future Scope",
        "• Physical microcontroller relay integration (Arduino / Raspberry Pi).",
        "• Multi-intersection cooperative reinforcement learning (MARL).",
        "• Emergency vehicle acoustic & visual priority preemption.",
        "• Night-mode support via infrared / thermal camera feeds.",
        "• Edge deployment on embedded GPU (Jetson Nano).",
    ]
    set_tf_lines(tf, future_lines, start_para=0)

if len(content_placeholders7) >= 3:
    tf = content_placeholders7[2].text_frame
    lim_lines = [
        "Limitations",
        "• Image quality degrades in heavy fog / night without infrared sensors.",
        "• Signal actuation is simulation-based, not physically wired to hardware.",
        "• Trained on a single dataset; real-world generalization needs broader data.",
        "• Single-intersection scope — no multi-intersection cooperation yet.",
    ]
    set_tf_lines(tf, lim_lines, start_para=0)

for shape in slide7.shapes:
    if shape.has_text_frame and "Date" in shape.name:
        set_para_text(shape.text_frame.paragraphs[0], f"Date: {TODAY}")

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 8 — Thank You / GitHub
# ─────────────────────────────────────────────────────────────────────────────
slide8 = slides[7]
subtitles8 = find_subtitle_shapes(slide8)

if subtitles8:
    tf = subtitles8[0].text_frame
    lines8 = [
        "MITHUN PRASATH LOGESH, GIRISH.R",
        "Computer Science and Business Systems",
    ]
    for i, line in enumerate(lines8):
        if i < len(tf.paragraphs):
            set_para_text(tf.paragraphs[i], line)

# GitHub TextBox
for shape in slide8.shapes:
    if shape.has_text_frame and "TextBox" in shape.name:
        tf = shape.text_frame
        set_para_text(tf.paragraphs[0], "GitHub: https://github.com/MithunPrasath/Cognisignal")
        break

for shape in slide8.shapes:
    if shape.has_text_frame and "Date" in shape.name:
        set_para_text(shape.text_frame.paragraphs[0], f"Date: {TODAY}")

# ─────────────────────────────────────────────────────────────────────────────
# Save
# ─────────────────────────────────────────────────────────────────────────────
prs.save(OUT_PPTX)
print(f"Saved: {OUT_PPTX}")
